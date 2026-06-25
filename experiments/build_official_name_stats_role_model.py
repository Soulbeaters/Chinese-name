#!/usr/bin/env python3
"""Build a frozen token role model from official national name statistics."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import sys
import zipfile
from collections import defaultdict
from pathlib import Path
from typing import Iterable

import openpyxl


PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

from experiments.evaluate_mentor_large_scale import normalized_tokens  # noqa: E402


US_CENSUS_2020_FIRST_NAMES_URL = (
    "https://www2.census.gov/topics/genealogy/2020surnames/"
    "Names2020_FirstNames_Sex.xlsx"
)
US_CENSUS_2020_LAST_NAMES_URL = (
    "https://www2.census.gov/topics/genealogy/2020surnames/"
    "Names2020_LastNames_RaceHispanic.xlsx"
)
FR_INSEE_PRENOMS_URL = (
    "https://www.insee.fr/fr/statistiques/fichier/8595130/"
    "prenoms-2024-nat_csv.zip"
)
FR_INSEE_NOMS_URL = (
    "https://www.insee.fr/fr/statistiques/fichier/3536630/"
    "noms2008nat_txt.zip"
)
PL_PESEL_GIVEN_FEMALE_URL = (
    "https://api.dane.gov.pl/resources/1159670,"
    "lista-imion-zenskich-w-rejestrze-pesel-stan-na-20012026-imie-pierwsze/csv"
)
PL_PESEL_GIVEN_MALE_URL = (
    "https://api.dane.gov.pl/resources/1159669,"
    "lista-imion-meskich-w-rejestrze-pesel-stan-na-20012026-imie-pierwsze/csv"
)
PL_PESEL_SURNAMES_FEMALE_URL = (
    "https://api.dane.gov.pl/resources/1148811,"
    "nazwiska-zenskie-stan-na-2026-01-20/csv"
)
PL_PESEL_SURNAMES_MALE_URL = (
    "https://api.dane.gov.pl/resources/1148808,"
    "nazwiska-meskie-stan-na-2026-01-20/csv"
)
SE_SCB_NAMES_URL = (
    "https://doris.snd.se//api/file/2021-272-1/1/data?"
    "filePath=Statistik%20%C3%B6ver%20namn%20efter%20f%C3%B6delseland%202020.zip"
)

TOKEN_RE = re.compile(r"^[a-z][a-z'-]*$")


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def single_token(value: object) -> str | None:
    tokens = normalized_tokens(str(value or ""))
    if len(tokens) != 1:
        return None
    token = tokens[0]
    if len(token.replace("-", "").replace("'", "")) < 2:
        return None
    if not TOKEN_RE.fullmatch(token):
        return None
    return token


def add_count(counts: dict[str, int], name: object, count: object) -> None:
    token = single_token(name)
    if not token:
        return
    try:
        value = int(float(str(count).strip()))
    except (TypeError, ValueError):
        return
    if value > 0:
        counts[token] += value


def load_us_census_first_names(path: Path) -> dict[str, int]:
    counts: dict[str, int] = defaultdict(int)
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    for row in sheet.iter_rows(min_row=4, values_only=True):
        add_count(counts, row[0], row[2])
    return dict(counts)


def load_us_census_last_names(path: Path) -> dict[str, int]:
    counts: dict[str, int] = defaultdict(int)
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    for row in sheet.iter_rows(min_row=4, values_only=True):
        add_count(counts, row[0], row[2])
    return dict(counts)


def load_fr_insee_prenoms(path: Path) -> dict[str, int]:
    counts: dict[str, int] = defaultdict(int)
    with zipfile.ZipFile(path) as archive:
        with archive.open("prenoms-2024-nat.csv") as handle:
            text = (line.decode("utf-8") for line in handle)
            for row in csv.DictReader(text, delimiter=";"):
                add_count(counts, row.get("prenom"), row.get("valeur"))
    return dict(counts)


def load_fr_insee_noms(path: Path) -> dict[str, int]:
    counts: dict[str, int] = defaultdict(int)
    with zipfile.ZipFile(path) as archive:
        with archive.open("noms2008nat_txt.txt") as handle:
            text = (line.decode("utf-8") for line in handle)
            for row in csv.DictReader(text, delimiter="\t"):
                total = 0
                for key, value in row.items():
                    if key and key.startswith("_"):
                        try:
                            total += int(value)
                        except (TypeError, ValueError):
                            pass
                add_count(counts, row.get("NOM"), total)
    return dict(counts)


def load_polish_csv(paths: Iterable[Path], name_columns: tuple[str, ...]) -> dict[str, int]:
    counts: dict[str, int] = defaultdict(int)
    for path in paths:
        with path.open(encoding="utf-8-sig", newline="") as handle:
            for row in csv.DictReader(handle):
                name = next((row.get(column) for column in name_columns if row.get(column)), None)
                count = row.get("LICZBA_WYSTĄPIEŃ") or row.get("Liczba")
                add_count(counts, name, count)
    return dict(counts)


def load_swedish_zip(path: Path, members: Iterable[str]) -> dict[str, int]:
    counts: dict[str, int] = defaultdict(int)
    with zipfile.ZipFile(path) as archive:
        for member in members:
            with archive.open(member) as handle:
                text = (line.decode("utf-8") for line in handle)
                for row in csv.DictReader(text):
                    add_count(counts, row.get("name"), row.get("total"))
    return dict(counts)


def normalized_role_counts(
    country_sources: dict[str, tuple[dict[str, int], dict[str, int]]],
    scale: int,
) -> dict[str, list[int]]:
    aggregate: dict[str, list[int]] = defaultdict(lambda: [0, 0])
    for surname_counts, given_counts in country_sources.values():
        surname_total = sum(surname_counts.values())
        given_total = sum(given_counts.values())
        if surname_total <= 0 or given_total <= 0:
            continue
        for token, count in surname_counts.items():
            aggregate[token][0] += max(1, round(count / surname_total * scale))
        for token, count in given_counts.items():
            aggregate[token][1] += max(1, round(count / given_total * scale))
    return dict(sorted(aggregate.items()))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--us-census-first", type=Path, required=True)
    parser.add_argument("--us-census-last", type=Path, required=True)
    parser.add_argument("--fr-insee-prenoms", type=Path, required=True)
    parser.add_argument("--fr-insee-noms", type=Path, required=True)
    parser.add_argument("--pl-given-female", type=Path, required=True)
    parser.add_argument("--pl-given-male", type=Path, required=True)
    parser.add_argument("--pl-surnames-female", type=Path, required=True)
    parser.add_argument("--pl-surnames-male", type=Path, required=True)
    parser.add_argument("--se-scb-zip", type=Path)
    parser.add_argument("--scale", type=int, default=1_000_000)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    us_given = load_us_census_first_names(args.us_census_first)
    us_surname = load_us_census_last_names(args.us_census_last)
    fr_given = load_fr_insee_prenoms(args.fr_insee_prenoms)
    fr_surname = load_fr_insee_noms(args.fr_insee_noms)
    pl_given = load_polish_csv(
        [args.pl_given_female, args.pl_given_male],
        ("IMIĘ_PIERWSZE",),
    )
    pl_surname = load_polish_csv(
        [args.pl_surnames_female, args.pl_surnames_male],
        ("Nazwisko aktualne",),
    )

    country_sources = {
        "us_census_2020": (us_surname, us_given),
        "fr_insee": (fr_surname, fr_given),
        "pl_pesel": (pl_surname, pl_given),
    }
    if args.se_scb_zip:
        se_surname = load_swedish_zip(args.se_scb_zip, ["lastname_wide.csv"])
        se_given = load_swedish_zip(args.se_scb_zip, ["men_wide.csv", "women_wide.csv"])
        country_sources["se_scb_2020"] = (se_surname, se_given)

    counts = normalized_role_counts(country_sources, args.scale)
    payload = {
        "version": 1,
        "description": (
            "Frozen token role counts from official national name statistics. "
            "Counts are source-normalized before aggregation."
        ),
        "scale_per_source_role": args.scale,
        "sources": {
            "us_census_2020_first_names_by_sex": {
                "url": US_CENSUS_2020_FIRST_NAMES_URL,
                "sha256": sha256_file(args.us_census_first),
            },
            "us_census_2020_last_names_by_race_hispanic": {
                "url": US_CENSUS_2020_LAST_NAMES_URL,
                "sha256": sha256_file(args.us_census_last),
            },
            "fr_insee_prenoms_2024_national": {
                "url": FR_INSEE_PRENOMS_URL,
                "sha256": sha256_file(args.fr_insee_prenoms),
            },
            "fr_insee_noms_2008_national": {
                "url": FR_INSEE_NOMS_URL,
                "sha256": sha256_file(args.fr_insee_noms),
            },
            "pl_pesel_given_female_2026_first": {
                "url": PL_PESEL_GIVEN_FEMALE_URL,
                "sha256": sha256_file(args.pl_given_female),
            },
            "pl_pesel_given_male_2026_first": {
                "url": PL_PESEL_GIVEN_MALE_URL,
                "sha256": sha256_file(args.pl_given_male),
            },
            "pl_pesel_surnames_female_2026": {
                "url": PL_PESEL_SURNAMES_FEMALE_URL,
                "sha256": sha256_file(args.pl_surnames_female),
            },
            "pl_pesel_surnames_male_2026": {
                "url": PL_PESEL_SURNAMES_MALE_URL,
                "sha256": sha256_file(args.pl_surnames_male),
            },
            **(
                {
                    "se_scb_names_by_birth_country_2020": {
                        "url": SE_SCB_NAMES_URL,
                        "landing_page": "https://researchdata.se/en/catalogue/dataset/2021-272-1",
                        "sha256": sha256_file(args.se_scb_zip),
                    }
                }
                if args.se_scb_zip
                else {}
            ),
        },
        "raw_totals": {
            country: {
                "surname": sum(surname_counts.values()),
                "given": sum(given_counts.values()),
                "surname_tokens": len(surname_counts),
                "given_tokens": len(given_counts),
            }
            for country, (surname_counts, given_counts) in country_sources.items()
        },
        "totals": {
            "surname": sum(value[0] for value in counts.values()),
            "given": sum(value[1] for value in counts.values()),
            "vocabulary": len(counts),
        },
        "counts": counts,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps(payload, ensure_ascii=False, separators=(",", ":")),
        encoding="utf-8",
    )
    print(
        json.dumps(
            {
                "output": str(args.output),
                "raw_totals": payload["raw_totals"],
                "totals": payload["totals"],
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
